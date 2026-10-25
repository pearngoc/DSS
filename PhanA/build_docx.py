# -*- coding: utf-8 -*-
"""Dựng báo cáo Word Phần A (Data-Driven DSS) từ kpi_summary.json + star_schema.png."""
import json
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
S = json.load(open(os.path.join(HERE, "tableau_data", "kpi_summary.json")))
NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)


def money(x):
    return "${:,.0f}".format(x)


doc = Document()
# font mặc định
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def h(text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = NAVY
    return p


def para(text, italic=False, bold=False, size=11, color=None, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic; r.bold = bold; r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        r = p.add_run(bold_prefix); r.bold = True
    p.add_run(text)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(hd); run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); run.font.size = Pt(10)
        _shade(c, "1F4E79")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val)); run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t

# ============================================================ TRANG BÌA
p = para("CAO HỌC K34 – HỆ HỖ TRỢ RA QUYẾT ĐỊNH", bold=True, size=13, color=GREY,
         align=WD_ALIGN_PARAGRAPH.CENTER)
para("ĐỒ ÁN DSS – PHẦN A: DATA-DRIVEN DSS", bold=True, size=20, color=NAVY,
     align=WD_ALIGN_PARAGRAPH.CENTER)
para("Hệ hỗ trợ ra quyết định phân tích Bán hàng & Giao hàng Chuỗi cung ứng toàn cầu "
     "(Bộ dữ liệu DataCo Supply Chain)", italic=True, size=12,
     align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
table(["Hạng mục", "Nội dung"],
      [["Môn học", "Hệ hỗ trợ ra quyết định (DSS)"],
       ["Đề tài", "Data-Driven DSS cho chuỗi cung ứng DataCo"],
       ["Công cụ Phần A", "Tableau (trực quan hoá) + Python (ETL/làm sạch dữ liệu)"],
       ["Mã nhóm", "<MaNhom> — điền trước khi nộp"],
       ["Thành viên & phân công", "<MSHV1 – …>, <MSHV2 – …>, <MSHV3 – …>"]],
      widths=[1.8, 4.5])
para("Ghi chú: các trường <…> là chỗ dành cho nhóm điền thông tin.", italic=True,
     size=9, color=GREY)
doc.add_page_break()

# ============================================================ 1. CHỦ ĐỀ
h("1. Giới thiệu chủ đề & nhu cầu phân tích", 1)
para("Chủ đề nhóm lựa chọn là quản trị chuỗi cung ứng – thương mại điện tử toàn cầu, "
     "dựa trên bộ dữ liệu giao dịch của công ty DataCo Global. Doanh nghiệp bán sản phẩm "
     "thể thao – thời trang – công nghệ qua nhiều thị trường (Market) trên thế giới và "
     "đối mặt hai nhóm vấn đề cốt lõi: (1) hiệu quả bán hàng & lợi nhuận, (2) chất lượng "
     "giao hàng (đúng hẹn / trễ hẹn). Đây là bài toán điển hình mà một Hệ hỗ trợ ra quyết "
     "định hướng dữ liệu (Data-Driven DSS) có thể giải quyết thông qua phân tích mô tả và "
     "trực quan hoá.")
para("Các nhu cầu phân tích đặt ra:")
bullet("doanh thu, lợi nhuận và biên lợi nhuận đang ở mức nào, theo thị trường/ngành hàng/thời gian?")
bullet("tỷ lệ giao hàng trễ bao nhiêu, tập trung ở phương thức vận chuyển hay khu vực nào?")
bullet("những đơn hàng/ngành hàng nào đang bị lỗ, nguyên nhân có phải do chiết khấu?")
bullet("nhóm khách hàng và thị trường nào tạo ra giá trị lớn nhất để ưu tiên nguồn lực?")

# ============================================================ 2. DATASET
h("2. Giới thiệu dataset", 1)
bullet("Kaggle – DataCo Smart Supply Chain for Big Data Analysis.", "Nguồn: ")
bullet("DataCoSupplyChainDataset.csv (mã hoá latin-1).", "Tệp: ")
bullet(f"{S['raw_rows']:,} dòng × {S['raw_cols']} cột ở dạng thô.", "Quy mô thô: ")
bullet("mỗi dòng là một mặt hàng trong đơn hàng (order item line-item).", "Mức chi tiết (grain): ")
bullet(f"{S['date_min']} đến {S['date_max']} (hơn 3 năm giao dịch).", "Khoảng thời gian: ")
para("Các trường dữ liệu chính (vượt yêu cầu tối thiểu 6 trường):")
table(["Nhóm", "Trường tiêu biểu"],
      [["Giao dịch/đo lường", "Sales, Order Item Total, Order Profit Per Order, Order Item Quantity, Order Item Discount Rate"],
       ["Khách hàng", "Customer Id, Customer Segment, Customer City/State/Country"],
       ["Sản phẩm", "Product Name, Category Name, Department Name, Product Price"],
       ["Thời gian", "order date (DateOrders), shipping date (DateOrders)"],
       ["Địa lý đơn hàng", "Market, Order Region, Order Country/State/City"],
       ["Giao hàng", "Shipping Mode, Delivery Status, Days for shipping (real/scheduled), Late_delivery_risk"]],
      widths=[1.8, 4.5])

# ============================================================ 3. LÀM SẠCH
h("3. Xử lý & làm sạch dữ liệu", 1)
para("Toàn bộ quy trình được tự động hoá bằng script Python (pandas) trong tệp "
     "etl_phanA.py. Các bước làm sạch:")
para("3.1. Loại bỏ cột không mang thông tin phân tích", bold=True)
table(["Cột bị loại", "Lý do"],
      [["Product Description", "trống 100%"],
       ["Order Zipcode", "trống 86,2%"],
       ["Customer Email, Customer Password", "chỉ 1 giá trị giả (ẩn danh), vô nghĩa"],
       ["Product Status", "toàn bộ = 0 (hằng số)"],
       ["Product Image", "đường dẫn ảnh, không dùng phân tích"],
       ["Benefit per order", "trùng lặp hoàn toàn với Order Profit Per Order"]],
      widths=[2.6, 3.7])
para("3.2. Xử lý giá trị trống", bold=True)
bullet("Customer Lname: 8 dòng trống → điền \"Unknown\".")
bullet("Customer Zipcode: 3 dòng trống → điền 0.")
para("3.3. Chuẩn hoá định dạng & trùng lặp", bold=True)
bullet(f"Kiểm tra trùng lặp toàn dòng: {S['dup_removed']} dòng trùng (dữ liệu sạch sẵn).")
bullet("Chuyển order date / shipping date (chuỗi M/D/YYYY) sang kiểu datetime chuẩn.")
bullet("Trim khoảng trắng thừa ở các cột văn bản (ví dụ \"Health and Beauty \").")
para("3.4. Tạo trường dẫn xuất", bold=True)
bullet("shipping_delay = Days for shipping (real) − Days for shipment (scheduled): số ngày "
       "giao trễ so với kế hoạch, phục vụ KPI độ trễ.")
para(f"Kết quả: dữ liệu sạch còn {S['clean_rows']:,} dòng, đồng nhất và sẵn sàng đưa vào "
     "mô hình phân tích.")

# ============================================================ 4. MỤC TIÊU DASHBOARD
h("4. Mục tiêu dashboard & đối tượng sử dụng", 1)
bullet("giám sát đồng thời hiệu quả BÁN HÀNG (doanh thu, lợi nhuận, biên lợi nhuận) và "
       "chất lượng GIAO HÀNG (tỷ lệ trễ, độ trễ) của chuỗi cung ứng toàn cầu, giúp phát "
       "hiện nhanh khu vực/ngành hàng có vấn đề và ra quyết định điều chỉnh.", "Mục tiêu dashboard: ")
bullet("Giám đốc Vận hành / Chuỗi cung ứng (theo dõi giao hàng), Giám đốc Kinh doanh và "
       "quản lý vùng (theo dõi doanh thu – lợi nhuận theo thị trường).", "Đối tượng sử dụng: ")

# ============================================================ 5. KPI
h("5. Các KPI chính", 1)
para("Dashboard thể hiện 7 KPI (vượt mức tối thiểu 5). Cột \"Calculated field (Tableau)\" "
     "là công thức để tạo trực tiếp trong Tableau.")
kpi_rows = [
    ["1. Tổng doanh thu", "Tổng giá trị bán ra", "SUM([sales])",
     money(S['total_sales']), "Quy mô kinh doanh; nền để so sánh theo chiều."],
    ["2. Tổng lợi nhuận", "Tổng lãi/lỗ thực tế", "SUM([profit])",
     money(S['total_profit']), "Hiệu quả tài chính thực sự của hoạt động bán hàng."],
    ["3. Biên lợi nhuận %", "Lợi nhuận / Doanh thu", "SUM([profit]) / SUM([sales])",
     f"{S['profit_margin_pct']}%", "Cảnh báo ngành hàng/thị trường bào mòn lợi nhuận."],
    ["4. Tỷ lệ giao trễ %", "% đơn hàng giao trễ", "SUM([late_delivery_flag])/COUNT([order_id])",
     f"{S['late_delivery_rate_order_pct']}%", "KPI vận hành cốt lõi về chất lượng giao hàng."],
    ["5. Độ trễ giao TB (ngày)", "Ngày thực − ngày kế hoạch (cấp đơn)", "AVG([shipping_delay])",
     f"{S['avg_shipping_delay']} ngày", "Mức độ chậm so với cam kết để cải thiện SLA."],
    ["6. Tỷ lệ đơn lỗ %", "% dòng đơn có profit < 0",
     "SUM([Loss Flag])/COUNT([order_item_id]); Loss Flag = IF [profit]<0 THEN 1 ELSE 0 END",
     f"{S['loss_order_line_pct']}%", "Khoanh vùng giao dịch lỗ để xử lý chiết khấu/giá."],
    ["7. Giá trị đơn TB (AOV)", "Doanh thu ròng / số đơn",
     "SUM([net_sales])/COUNTD([order_id])",
     money(S['aov']), "Đo sức mua mỗi đơn; định hướng up-sell/cross-sell."],
]
table(["KPI", "Cách tính", "Calculated field (Tableau)", "Giá trị", "Ý nghĩa ra quyết định"],
      kpi_rows, widths=[1.1, 1.1, 1.6, 0.8, 1.7])

# ============================================================ 6. STAR SCHEMA
h("6. Thiết kế lược đồ sao (Star Schema)", 1)
para("Mục này trình bày thiết kế kho dữ liệu theo quy trình chuẩn của Kimball: xác định "
     "business process → grain → dimensions → facts/measures → loại fact table → vẽ lược "
     "đồ sao → kiểm chứng. Thiết kế gồm đúng hai bảng fact.")

# ---- 6.1 Business Process ----
h("6.1. Xác định Business Process", 2)
para("Từ đặc tả nghiệp vụ, dữ liệu DataCo chứa hai quy trình nghiệp vụ (business process) "
     "có mức chi tiết khác nhau, tương ứng hai bảng fact:")
bullet("phân tích bán hàng – lợi nhuận ở mức TỪNG MẶT HÀNG trong đơn (doanh thu, chiết "
       "khấu, lợi nhuận theo sản phẩm/khách hàng/thị trường).", "FACT_SALES — ")
bullet("phân tích hiệu quả giao hàng ở mức TỪNG ĐƠN HÀNG/LẦN VẬN CHUYỂN (thời gian giao, "
       "trễ hẹn, đúng hạn theo phương thức/khu vực).", "FACT_ORDER_DELIVERY — ")
para("Vì sao hai quy trình có grain khác nhau và KHÔNG gộp chung một fact table?", bold=True)
bullet("Bán hàng phát sinh ở mức dòng mặt hàng (một đơn nhiều sản phẩm, mỗi sản phẩm một "
       "doanh thu/lợi nhuận riêng). Giao hàng lại là thuộc tính của CẢ ĐƠN: ngày giao, số "
       "ngày vận chuyển, trạng thái trễ là như nhau cho mọi mặt hàng trong đơn.")
bullet(f"Nếu nhét chỉ số giao hàng vào fact mức mặt hàng, chúng sẽ bị lặp lại trên mọi dòng "
       f"(trung bình {S['avg_items_per_order']} dòng/đơn) → khi tính 'tỷ lệ giao trễ' sẽ "
       "ĐẾM TRÙNG theo số mặt hàng, thiên lệch về các đơn lớn. Tách fact giúp mỗi đơn chỉ "
       "được tính một lần.")
bullet("Trộn hai grain trong một bảng vi phạm nguyên tắc 'một grain cho một fact' và làm "
       "các phép tổng hợp (SUM, AVG) không còn chính xác.")

# ---- 6.2 Grain ----
h("6.2. Xác định Grain", 2)
bullet("“Một dòng trong FACT_SALES biểu diễn một mặt hàng (line-item) thuộc một đơn "
       "hàng.”", "")
bullet("“Một dòng trong FACT_ORDER_DELIVERY biểu diễn một đơn hàng được vận chuyển "
       "tới người nhận.”", "")
para("Kiểm tra grain có đáp ứng yêu cầu phân tích:")
table(["Yêu cầu", "Grain hỗ trợ?", "Giải thích"],
      [["KPI doanh thu/lợi nhuận theo sản phẩm", "Có", "Grain mặt hàng của FACT_SALES đủ chi tiết để roll-up lên category, department."],
       ["KPI tỷ lệ giao trễ / đúng hạn", "Có", "Grain đơn hàng của FACT_ORDER_DELIVERY tính đúng mỗi đơn một lần."],
       ["Drill-down Year→Quarter→Month→Date", "Có", "Cả hai fact đều có khoá ngày (order_date_key, ship_date_key)."],
       ["Drill-down Market→Region→Country→State→City", "Có", "Cả hai fact đều nối DIM_LOCATION."]],
      widths=[2.3, 0.9, 3.0])

# ---- 6.3 Dimensions ----
h("6.3. Xác định Dimensions", 2)
para("Các dimension cùng business key (BK), surrogate key (SK), thuộc tính mô tả và phân cấp:")
table(["Dimension", "Business key", "Surrogate key", "Thuộc tính mô tả", "Hierarchy", "Số dòng"],
      [["DIM_DATE", "full_date", "date_key (YYYYMMDD)", "year, quarter, month, weekday, is_weekend", "Year→Quarter→Month→Date", f"{S['tables']['dim_date'][0]:,}"],
       ["DIM_CUSTOMER", "Customer Id", "customer_key", "segment, city, state, country", "Country→State→City", f"{S['tables']['dim_customer'][0]:,}"],
       ["DIM_PRODUCT", "Product Card Id", "product_key", "product_name, category, department, product_price", "Department→Category→Product", f"{S['tables']['dim_product'][0]:,}"],
       ["DIM_LOCATION", "(market..city)", "location_key", "market, region, country, state, city", "Market→Region→Country→State→City", f"{S['tables']['dim_location'][0]:,}"],
       ["DIM_DELIVERY", "(mode..status)", "delivery_key", "shipping_mode, delivery_status, payment_type, order_status", "—", f"{S['tables']['dim_delivery'][0]:,}"]],
      widths=[1.2, 1.0, 1.1, 1.6, 1.4, 0.6])
para("Conformed dimensions (chia sẻ giữa hai fact): ", bold=True)
bullet("DIM_DATE, DIM_CUSTOMER, DIM_LOCATION, DIM_DELIVERY dùng chung cho cả FACT_SALES và "
       "FACT_ORDER_DELIVERY — cho phép so sánh/kết hợp hai nghiệp vụ trên cùng bộ chiều.")
bullet("DIM_PRODUCT chỉ áp dụng cho FACT_SALES (mức đơn hàng không gắn với một sản phẩm "
       "duy nhất).")
para("Role-playing dimension: ", bold=True)
bullet("DIM_DATE đóng hai vai trò trong cùng một fact qua hai khoá order_date_key (ngày "
       "đặt) và ship_date_key (ngày giao). Trong Tableau sẽ nối DIM_DATE hai lần.")

# ---- 6.4 Facts & Measures ----
h("6.4. Xác định Facts và Measures", 2)
para("FACT_SALES — base measures (lưu tại grain mặt hàng):", bold=True)
table(["Measure", "Ý nghĩa", "Cách tính", "Đơn vị", "Khả năng tổng hợp"],
      [["quantity", "Số lượng bán", "nguồn", "cái", "Additive"],
       ["sales", "Doanh thu gộp", "quantity × unit_price", "USD", "Additive"],
       ["discount", "Tiền chiết khấu", "nguồn", "USD", "Additive"],
       ["net_sales", "Doanh thu ròng", "sales − discount", "USD", "Additive"],
       ["profit", "Lợi nhuận", "nguồn (Order Profit Per Order)", "USD", "Additive"],
       ["unit_price", "Đơn giá", "nguồn", "USD", "Non-additive (giá)"],
       ["profit_ratio", "Tỷ suất LN dòng", "profit / net_sales", "%", "Non-additive (tỷ lệ)"]],
      widths=[1.0, 1.3, 1.6, 0.7, 1.4])
para("FACT_ORDER_DELIVERY — base measures (lưu tại grain đơn hàng):", bold=True)
table(["Measure", "Ý nghĩa", "Cách tính", "Đơn vị", "Khả năng tổng hợp"],
      [["days_ship_scheduled", "Số ngày giao cam kết", "nguồn", "ngày", "Semi-additive (TB theo đơn)"],
       ["days_ship_real", "Số ngày giao thực tế", "nguồn", "ngày", "Semi-additive"],
       ["shipping_delay", "Độ trễ", "real − scheduled", "ngày", "Semi-additive"],
       ["late_delivery_flag", "Cờ giao trễ (0/1)", "nguồn Late_delivery_risk", "cờ", "Additive (đếm số trễ)"],
       ["on_time_flag", "Cờ đúng hạn (0/1)", "1 − late_delivery_flag", "cờ", "Additive"],
       ["order_sales", "Doanh thu đơn", "Σ sales trong đơn", "USD", "Additive"],
       ["order_profit", "Lợi nhuận đơn", "Σ profit trong đơn", "USD", "Additive"]],
      widths=[1.2, 1.2, 1.3, 0.7, 1.5])
para("Base measures vs Calculated measures/KPI:", bold=True)
bullet("chỉ lưu các đại lượng CÓ THỂ CỘNG (quantity, sales, profit, late_delivery_flag…) "
       "trong fact table.", "Lưu trực tiếp (base): ")
bullet("các KPI dạng tỷ lệ/trung bình KHÔNG lưu sẵn mà tính ở tầng OLAP/BI từ base measures, "
       "ví dụ Biên LN% = SUM(profit)/SUM(sales); Tỷ lệ giao trễ% = SUM(late_delivery_flag)/"
       "COUNT(order_id); AOV = SUM(order_net_sales)/COUNTD(order_id). Lý do: tỷ lệ không "
       "cộng được nên nếu lưu sẵn sẽ sai khi roll-up.", "Tính ở OLAP (calculated): ")
para("Kiểm chứng khả năng tổng hợp: các base measures cộng được bảo đảm tính lại đúng mọi "
     "KPI khi roll-up/drill-down/slice/dice — ví dụ tỷ lệ giao trễ ở bất kỳ mức nào (thị "
     "trường, tháng) đều = Σ(late_delivery_flag)/Σ(count) tại mức đó, luôn chính xác.",
     italic=True, size=9, color=GREY)

# ---- 6.5 Loại Fact Table ----
h("6.5. Xác định loại Fact Table", 2)
table(["Fact table", "Loại", "Lý do"],
      [["FACT_SALES", "Transaction fact", "Ghi nhận từng giao dịch bán (mặt hàng) tại thời điểm phát sinh; grain nguyên tử."],
       ["FACT_ORDER_DELIVERY", "Accumulating snapshot", "Theo vòng đời đơn hàng với các mốc ngày đặt → ngày giao và các đại lượng độ trễ tích luỹ giữa các mốc."]],
      widths=[2.0, 1.6, 2.7])

# ---- 6.6 Sơ đồ ----
h("6.6. Lược đồ Star Schema", 2)
doc.add_picture(os.path.join(HERE, "star_schema.png"), width=Inches(6.5))
doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
para("Hình 1. Lược đồ sao hai fact của DSS chuỗi cung ứng DataCo (thể hiện PK, FK, surrogate "
     "key, measures, quan hệ, grain, conformed & role-playing dimension).",
     italic=True, size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

# ---- 6.7 Kiểm chứng ----
h("6.7. Kiểm chứng thiết kế", 2)
para("Mô hình trả lời được các câu hỏi phân tích yêu cầu (đã điều chỉnh theo dữ liệu DataCo — "
     "bộ dữ liệu không có thông tin nhân viên giao hàng, nhiều lần giao hay lý do thất bại "
     "chi tiết, nên các câu hỏi liên quan được thay bằng chỉ số tương đương gần nhất):")
table(["Câu hỏi phân tích", "Fact dùng", "Dimension / drill", "Measure / KPI"],
      [["1. Quý gần nhất, khu vực nào tỷ lệ đúng hạn thấp nhất?", "Order_Delivery", "Dim_Date (lọc quý), Dim_Location→region", "AVG(on_time_flag)"],
       ["2. Phương thức nào giao chậm hơn cam kết nhiều nhất?", "Order_Delivery", "Dim_Delivery→shipping_mode", "AVG(shipping_delay)"],
       ["3. Tỷ lệ giao đúng hạn theo khu vực (thay cho 'thành công lần đầu')?", "Order_Delivery", "Dim_Location→region", "AVG(on_time_flag)"],
       ["4. Yếu tố nào gây trễ nhiều nhất (thay cho 'lý do thất bại')?", "Order_Delivery", "Dim_Delivery (mode, status)", "SUM(late_delivery_flag)"],
       ["5. Khu vực nào tỷ lệ giao trễ cao nhất (thay cho 'nhân viên')?", "Order_Delivery", "Dim_Location→region/state", "AVG(late_delivery_flag)"],
       ["6. Doanh thu Year→Quarter→Month theo phương thức?", "Sales", "Dim_Date (drill), Dim_Delivery→shipping_mode", "SUM(sales)"],
       ["7. Nhóm khách hàng nào doanh thu lớn nhất?", "Sales", "Dim_Customer→segment", "SUM(sales)"],
       ["8. Khu vực nào nhiều đơn nhưng đúng hạn thấp?", "Order_Delivery", "Dim_Location→region", "COUNT(order_id) vs AVG(on_time_flag)"]],
      widths=[2.4, 1.0, 1.9, 1.3])
para("Mọi câu hỏi đều drill-down được tới mức yêu cầu nhờ phân cấp của DIM_DATE và "
     "DIM_LOCATION; grain, dimensions và measures do đó đã đủ, không cần điều chỉnh.",
     italic=True, size=9, color=GREY)

# ============================================================ 7. ETL
h("7. Truy vấn dữ liệu & ETL vào Star Schema", 1)
para("Quy trình ETL (Extract – Transform – Load) được thực hiện trong etl_phanA.py:")
bullet("đọc CSV thô DataCoSupplyChainDataset.csv.", "Extract: ")
bullet("làm sạch (mục 3), tách thuộc tính thành các bảng dimension (drop_duplicates theo khoá), "
       "tạo DIM_DATE từ hợp của ngày đặt & ngày giao, sinh khoá thay thế (surrogate key) cho "
       "DIM_LOCATION và DIM_DELIVERY; dựng FACT_ORDER_DELIVERY bằng cách gộp FACT_SALES về "
       "grain đơn hàng (group by order_id).", "Transform: ")
bullet(f"gắn khoá ngoại vào các bảng fact bằng phép merge, xuất {len(S['tables'])} tệp CSV.", "Load: ")
para("Các tệp CSV kết quả (thư mục tableau_data/) — import trực tiếp vào Tableau:")
table(["Tệp CSV", "Vai trò", "Kích thước (dòng × cột)"],
      [[f"{k}.csv", ("Fact" if k.startswith("fact") else "Dimension"),
        f"{v[0]:,} × {v[1]}"] for k, v in S['tables'].items()],
      widths=[1.9, 1.4, 2.0])
para("Ví dụ truy vấn kiểm chứng dữ liệu đã tổ chức hợp lý (doanh thu & tỷ lệ trễ theo thị trường):",
     size=10)
cq = doc.add_paragraph()
rq = cq.add_run("SELECT l.market,\n"
                "       SUM(d.order_sales)        AS doanh_thu,\n"
                "       AVG(d.late_delivery_flag) AS ty_le_tre\n"
                "FROM   fact_order_delivery d\n"
                "JOIN   dim_location l ON d.location_key = l.location_key\n"
                "GROUP  BY l.market\n"
                "ORDER  BY doanh_thu DESC;")
rq.font.name = "Consolas"; rq.font.size = Pt(9)

# ============================================================ 8. TABLEAU
h("8. Hướng dẫn thiết kế dashboard trên Tableau", 1)
para("8.1. Ứng dụng sử dụng & kết nối dữ liệu", bold=True)
bullet("dùng Tableau Desktop (bản Student miễn phí 1 năm qua Tableau for Students) hoặc Tableau "
       "Public (miễn phí; workbook được lưu công khai trên Tableau Public cloud — chấp nhận được "
       "vì DataCo là dữ liệu công khai trên Kaggle). Không cần Tableau Prep vì ETL đã làm bằng "
       "Python; không cần Power BI.", "Ứng dụng: ")
bullet("Connect → Text file → chọn các CSV trong tableau_data/. Kiểm tra kiểu dữ liệu: "
       "full_date = Date; các cột *_key, order_id, order_item_id, year, month = Dimension "
       "(chuột phải → Convert to Dimension) để Tableau không cộng dồn chúng như measure.",
       "Kết nối: ")
bullet("tạo HAI data source vì hai fact khác grain (nếu nối chung sẽ đếm trùng ~2,75 lần/đơn):",
       "Mô hình: ")
bullet("fact_sales ở giữa, Relationship tới dim_customer, dim_product, dim_location, "
       "dim_delivery (khớp theo *_key); dim_date nối hai lần: order_date_key và ship_date_key "
       "— đổi tên hai bảng logic thành Order Date và Ship Date (role-playing).",
       "Data source 1 « DataCo Sales »: ")
bullet("fact_order_delivery ở giữa, Relationship tới dim_customer, dim_location, "
       "dim_delivery, và dim_date (order_date_key; thêm lần hai cho ship_date_key nếu cần).",
       "Data source 2 « DataCo Delivery »: ")
bullet("Tableau không cho nối hai fact vào chung một mô hình khi chúng cùng chia sẻ nhiều "
       "dimension (sẽ tạo vòng lặp), nên cách tách hai data source là đúng. Hai nguồn dùng "
       "chung tên trường (market, shipping_mode, segment, year) để các Filter Action trên "
       "dashboard lọc xuyên hai nguồn.", "Lưu ý: ")
bullet("khi kéo order_date_key và ship_date_key cùng vào một quan hệ với dim_date, Tableau tạo "
       "điều kiện AND (cả hai khóa phải bằng date_key), nên chỉ giữ các đơn giao cùng ngày đặt "
       "(5.080/180.519 dòng; doanh thu chỉ ~$1,03 triệu thay vì $36,78 triệu) và sinh giá trị Null "
       "ở trường ngày. Cách đúng: mỗi quan hệ chỉ dùng MỘT khóa. Với dashboard này, nối dim_date "
       "theo order_date_key; chỉ thêm một bản dim_date thứ hai cho ship_date_key khi cần phân tích "
       "theo ngày giao.", "Lỗi thường gặp: ")

para("8.2. Dựng biểu đồ chi tiết trong Tableau", bold=True)
para("Bước 1 — Tạo Calculated Field (Analysis → Create Calculated Field), sau đó chuột phải "
     "vào field → Default Properties → Number Format để đặt Percentage (1 chữ số thập phân) "
     "hoặc Currency ($, 0 chữ số thập phân):")
table(["Tên field", "Data source", "Công thức"],
      [["Profit Margin %", "Sales", "SUM([profit]) / SUM([sales])"],
       ["Loss Flag", "Sales", "IF [profit] < 0 THEN 1 ELSE 0 END"],
       ["Loss Rate %", "Sales", "SUM([Loss Flag]) / COUNT([order_item_id])"],
       ["AOV", "Sales", "SUM([net_sales]) / COUNTD([order_id])"],
       ["Late Rate %", "Delivery", "SUM([late_delivery_flag]) / COUNT([order_id])"],
       ["Avg Delay (ngày)", "Delivery", "AVG([shipping_delay])"]],
      widths=[1.5, 1.1, 3.2])
para("")
para("Bước 2 — Dựng từng Sheet (Worksheet). Cách chung: kéo field từ khung Data bên trái vào "
     "Columns / Rows / Marks (Color, Size, Label, Tooltip), hoặc chọn loại biểu đồ ở Show Me; "
     "đặt tên sheet theo cột đầu tiên của bảng dưới:")
table(["Sheet (nguồn)", "Loại biểu đồ", "Kéo thả & thiết lập", "Định dạng"],
      [["K1–K7 Thẻ KPI (Sales/ Delivery)", "Text (BAN)",
        "Mỗi KPI một sheet: kéo measure vào Text; Fit → Entire View; ẩn tiêu đề cột "
        "(chuột phải → Hide Field Labels). K1 SUM(sales), K2 SUM(profit), K3 Profit Margin %, "
        "K4 Late Rate %, K5 Avg Delay, K6 Loss Rate %, K7 AOV.",
        "Font 28 đậm; tiêu đề sheet = tên KPI; Late Rate % và Loss Rate % tô cam đỏ (#D55E00)."],
       ["S1 Xu hướng thời gian (Sales)", "Dual-axis line",
        "Columns: full_date (Order Date) → chuột phải chọn MONTH loại liên tục (Continuous). "
        "Rows: SUM(sales). Kéo SUM(profit) thả vào mép phải trục → Dual Axis.",
        "Không Synchronize Axis (hai thang đo khác nhau); Profit màu cam. Ghi chú: 01/2018 là "
        "tháng dở dang."],
       ["S2 Doanh thu theo thị trường (Sales)", "Bar ngang",
        "Rows: market; Columns: SUM(sales). Color: Profit Margin %. Label: SUM(sales).",
        "Sort giảm dần; dùng dải màu tuần tự (sequential) xanh nhạt → xanh đậm #DEEBF7 → #08519C: biên LN càng cao càng đậm."],
       ["S3 Doanh thu & biên LN theo ngành (Sales)", "Bar ngang",
        "Rows: department_name; Columns: SUM(sales); Color: Profit Margin %; Label: SUM(sales).",
        "Sort giảm dần; tooltip thêm SUM(profit)."],
       ["S4 Tỷ lệ giao trễ theo phương thức (Delivery)", "Highlight bar",
        "Rows: shipping_mode; Columns: Late Rate %; Color: Late Rate %; Tooltip: Avg Delay, "
        "COUNTD(order_id). Analytics → Reference Line (Constant 0,548).",
        "Dải cam tuần tự #FDE0C5 → #D55E00, càng trễ càng đậm; sort giảm dần; nhãn %."],
       ["S5 Phân khúc khách hàng (Sales)", "Treemap",
        "Marks: Square; Size: SUM(sales); Color: segment; Label: segment + SUM(sales) hiển "
        "thị % of total (Quick Table Calculation → Percent of Total).",
        "Chỉ 3 phân khúc nên giữ màu phân biệt rõ."],
       ["S6 Tỷ lệ đơn lỗ theo ngành (Sales)", "Bar ngang",
        "Rows: department_name; Columns: Loss Rate %.",
        "Kết quả đồng đều ~18% ở mọi ngành → lỗ mang tính hệ thống (giá/chiết khấu), "
        "không tập trung ở một ngành."],
       ["S7 Top 10 sản phẩm (Sales)", "Bar ngang",
        "Rows: product_name; Columns: SUM(sales). Kéo product_name vào Filters → tab Top → "
        "By field: Top 10 by SUM(sales). Color: Profit Margin %.",
        "Thay thế bảng chi tiết trên dashboard chính."]],
      widths=[1.3, 0.9, 2.6, 1.3])
para("")
para("Về bảng chi tiết: dashboard tổng quan KHÔNG nên chứa bảng nhiều dòng — người ra quyết "
     "định cần đọc được trong vài giây, bảng làm rối và đẩy các biểu đồ ra khỏi màn hình. Vì "
     "vậy bản cũ « Bảng chi tiết (drill) » được thay bằng biểu đồ S7 Top 10 sản phẩm. Bảng chi "
     "tiết (product_name, department_name, SUM(sales), SUM(profit), Profit Margin %) vẫn tạo "
     "dưới dạng sheet D1 (Text table) và đặt ở một dashboard thứ hai « Chi tiết », mở bằng nút "
     "Navigation hoặc Dashboard Action loại Go to Sheet khi cần soi sâu hoặc demo OLAP.",
     italic=True)

para("8.3. Hướng dẫn chi tiết các thao tác OLAP", bold=True)
para("Chuẩn bị phân cấp (làm một lần trong Data source 1):", bold=True)
bullet("kéo full_date của Order Date vào Columns; Tableau tạo YEAR(full_date). Bấm dấu (+) "
       "trên pill để khoan sâu Year → Quarter → Month → Day; bấm (−) để cuộn lên.",
       "Thời gian: ")
bullet("trong khung Data, kéo region thả lên market → đặt tên « Geography » → kéo thêm "
       "country, state, city vào theo thứ tự.", "Địa lý: ")
bullet("kéo category_name thả lên department_name → đặt tên « Product Hierarchy » → thêm "
       "product_name.", "Sản phẩm: ")
para("Các thao tác (số liệu trong ngoặc để đối chiếu khi demo):", bold=True)
bullet("(1) Mở sheet S2. (2) Kéo market vào Filters → chọn Europe → OK. (3) Chuột phải pill "
       "trong Filters → Show Filter → chọn Single Value (dropdown). Kết quả: doanh thu "
       "$10,87 triệu, biên LN 10,76%, ngành lớn nhất Fan Shop $4,74 triệu. Cách khác: nhấp một "
       "thanh trên S2 rồi chọn biểu tượng Use as Filter. Slice = cố định MỘT chiều.", "Slice: ")
bullet("(1) Giữ lọc market = Europe. (2) Kéo thêm YEAR(full_date) và department_name vào "
       "Filters, chọn 2017 và Fan Shop. Kết quả: $1,91 triệu, biên LN 11,01%. Để xem hộp "
       "chéo nhiều chiều: tạo sheet mới, Rows = department_name, Columns = market, Marks = "
       "Square, Color = Profit Margin % (Highlight table). Dice = lọc đồng thời "
       "NHIỀU chiều.", "Dice: ")
bullet("(1) Trên S1 bấm (+) ở YEAR → QUARTER → MONTH → DAY; mỗi lần bấm mức chi tiết tăng một "
       "cấp. Đối chiếu tổng theo năm: 2015 $12,34 triệu; 2016 $12,30 triệu; 2017 $11,81 triệu "
       "(−4,0% so với 2016); 2018 chỉ có tháng 1 ($0,33 triệu). (2) Trên S2 dùng Geography: bấm "
       "(+) ở market để xuống region → country. Roll-up: bấm (−) trên pill để quay về cấp cao.",
       "Drill-down / Roll-up: ")
bullet("(1) Sheet mới (Pivot_Dept_Year): Rows = department_name; Columns = YEAR(full_date); "
       "kéo SUM(sales) vào Text → có bảng chéo (cross-tab). (2) Nhấn Swap Rows and Columns "
       "(Ctrl+W) để xoay: năm xuống hàng, ngành lên cột. (3) Kéo market từ Filters sang "
       "Columns để thêm chiều thứ hai và so sánh. Lưu ý đây là pivot theo nghĩa OLAP (xoay "
       "góc nhìn), khác lệnh Pivot trong Data Source (chuyển cột thành hàng) — không cần dùng.",
       "Pivot: ")
bullet("Dashboard → Actions → Add Action → Filter: Source = các sheet S2–S3, Target = tất cả "
       "sheet, Run action on Select. Bấm một thị trường thì cả KPI và biểu đồ (cả nguồn Sales "
       "lẫn Delivery) đổi theo vì dùng chung tên trường market. Thêm Highlight Action cho "
       "shipping_mode.", "Tương tác: ")
para("Kịch bản demo 3 phút: Slice Europe → Dice thêm 2017 và Fan Shop → Drill-down theo "
     "tháng thấy xu hướng → Pivot ngành × năm → bỏ lọc, bấm First Class trên S4 để cho thấy "
     "chỉ số giao trễ.", italic=True)

para("8.4. Bố cục dashboard", bold=True)
para("Tạo dashboard (Dashboard → New Dashboard), Size: Fixed 1300×800. Kéo sheet vào theo "
     "bố cục:")
table(["Vùng", "Nội dung"],
      [["Hàng 1", "Tiêu đề « DataCo — Dashboard Bán hàng & Giao hàng » + 7 thẻ KPI K1–K7 "
                  "(Horizontal container)"],
       ["Hàng 2 — trái (60%)", "S1 Xu hướng doanh thu & lợi nhuận theo tháng"],
       ["Hàng 2 — phải (40%)", "S2 Doanh thu theo thị trường"],
       ["Hàng 3", "S3 Ngành hàng | S4 Tỷ lệ giao trễ theo phương thức | S5 Phân khúc KH "
                  "hoặc S7 Top 10 sản phẩm"],
       ["Cạnh phải", "Bộ lọc: Year, Market, Segment, Shipping Mode (Show Filter)"],
       ["Dashboard 2 « Chi tiết »", "D1 bảng sản phẩm + bảng pivot ngành × năm; nút « Quay "
                                    "lại » (Navigation)"]],
      widths=[1.7, 4.1])
bullet("tối đa 2–3 màu chủ đạo; cam đỏ #D55E00 chỉ dành cho chỉ số cảnh báo (trễ, lỗ); bảng màu Okabe-Ito an toàn cho người mù màu; mỗi biểu đồ có "
       "tiêu đề nêu thông điệp.")

para("8.5. Kiểm tra mức độ đầy đủ của dashboard", bold=True)
table(["Yêu cầu đề bài", "Đáp ứng bởi", "Kết luận"],
      [["≥ 5 KPI trên dashboard", "7 thẻ KPI K1–K7 (Mục 5)", "Đủ"],
       ["KPI thể hiện bằng biểu đồ/ bảng phù hợp", "Mỗi KPI có thẻ và ít nhất một biểu đồ "
        "phân rã: doanh thu/LN (S1–S3), biên LN (S2–S3), giao trễ & độ trễ (S4), đơn lỗ (S6), "
        "AOV (K7, S7)", "Đủ"],
       ["OLAP: slice/dice, drill-down, pivot", "Mục 8.3: Filter market, Highlight table, "
        "hierarchy Date/Geography/Product, cross-tab ngành × năm", "Đủ"],
       ["≥ 3 insight có bằng chứng trên dashboard", "Insight 1 ← S4; Insight 2 ← K3, K6, S6; "
        "Insight 3 ← S2, S3", "Đủ"],
       ["Dễ đọc, phục vụ ra quyết định", "Trang chính chỉ biểu đồ, bảng chi tiết tách sang "
        "dashboard 2", "Đủ"]],
      widths=[1.8, 3.3, 0.7])
para("")
para("Hạn chế cần nêu khi trình bày: (a) biên lợi nhuận hầu như đồng đều giữa các thị trường "
     "và ngành (≈ 10–11%) nên màu theo biên LN ít phân biệt; điểm khác biệt thực sự nằm ở quy "
     "mô doanh thu và tỷ lệ trễ theo phương thức giao hàng; (b) dữ liệu 2018 chỉ có tháng 1; "
     "(c) bản đồ không dùng vì tên quốc gia (tiếng Tây Ban Nha) khó nhận diện địa lý.",
     italic=True)


# ============================================================ 9. INSIGHT
h("9. Phân tích & khuyến nghị", 1)
para("Dựa trên dashboard, nhóm rút ra các insight quan trọng:")
lr = S['late_rate_by_shipmode_pct']
bullet(f"Giao hàng trễ ở mức báo động {S['late_delivery_rate_order_pct']}% tổng số đơn hàng. "
       f"Nghịch lý: phương thức nhanh/đắt tiền lại trễ nhiều nhất — First Class {lr.get('First Class')}%, "
       f"Second Class {lr.get('Second Class')}%, trong khi Standard Class chỉ {lr.get('Standard Class')}%.",
       "Insight 1 — ")
bullet(f"{S['loss_order_line_pct']}% dòng đơn ({S['loss_order_lines']:,} dòng) bị LỖ dù doanh thu dương, "
       f"kéo biên lợi nhuận toàn công ty xuống chỉ {S['profit_margin_pct']}% "
       f"(chiết khấu trung bình {S['avg_discount_rate_pct']}%).", "Insight 2 — ")
dep = S['sales_by_department']; top_dep = list(dep.items())[0]
mk = S['sales_by_market']
bullet(f"Doanh thu tập trung cao: ngành {top_dep[0]} chiếm {money(top_dep[1])} (lớn nhất), "
       f"và hai thị trường Europe ({money(mk['Europe'])}) + LATAM ({money(mk['LATAM'])}) "
       "dẫn đầu toàn cầu.", "Insight 3 — ")
para("Khuyến nghị ra quyết định:")
bullet("rà soát lại nhà vận chuyển và quy trình của First/Second Class — ưu tiên cố định "
       "SLA cho hai phương thức này vì đang trễ nặng nhất.")
bullet("thiết lập ngưỡng chiết khấu tối đa và cảnh báo khi profit_ratio âm, nhằm cắt giảm "
       f"{S['loss_order_line_pct']}% dòng đơn lỗ.")
bullet("tập trung nguồn lực marketing/tồn kho vào ngành và thị trường dẫn đầu, đồng thời "
       "khảo sát vì sao các thị trường như Africa còn nhỏ để tìm dư địa tăng trưởng.")

# ============================================================ LƯU
out = os.path.join(HERE, "PhanA_BaoCao.docx")
doc.save(out)
print("saved", out)
